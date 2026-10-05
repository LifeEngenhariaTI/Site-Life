from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path('Relatorio_Seguranca_Canal_Email.docx')
NAVY, PALE, LINE = '07263D', 'EEF5F7', 'D9D9D9'

def border(cell):
    pr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders'); pr.append(b)
    for side in ('top','left','bottom','right'):
        x = OxmlElement(f'w:{side}'); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'6'); x.set(qn('w:color'),LINE); b.append(x)

def cell(cell, text, header=False, alternate=False):
    if header or alternate:
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), NAVY if header else PALE); cell._tc.get_or_add_tcPr().append(shade)
    p = cell.paragraphs[0]; p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text); r.bold = header; r.font.name = 'Aptos'; r.font.size = Pt(9); r._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos')
    if header: r.font.color.rgb = RGBColor(255,255,255)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; border(cell)

def table(doc, headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers): cell(t.rows[0].cells[i],h,header=True)
    for n,row in enumerate(rows):
        cells = t.add_row().cells
        for i,value in enumerate(row): cell(cells[i],value,alternate=n%2==1)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def text(doc, value):
    p=doc.add_paragraph(); p.paragraph_format.line_spacing=1.15; p.paragraph_format.space_after=Pt(8)
    r=p.add_run(value); r.font.name='Aptos'; r.font.size=Pt(10.5); r._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos')

def heading(doc, value, level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); p.paragraph_format.space_before=Pt(13); p.paragraph_format.space_after=Pt(6)
    r=p.add_run(value); r.font.color.rgb=RGBColor(0,0,0)

doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.75); sec.bottom_margin=Inches(.75); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
for style in ('Normal','Title','Heading 1','Heading 2'):
    doc.styles[style].font.name='Aptos'; doc.styles[style]._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); doc.styles[style].font.color.rgb=RGBColor(0,0,0)
doc.styles['Title'].font.size=Pt(24); doc.styles['Heading 1'].font.size=Pt(15); doc.styles['Heading 2'].font.size=Pt(12)

p=doc.add_paragraph(style='Title'); p.add_run('Relatorio de seguranca do canal de email')
p=doc.add_paragraph('Formulario de transparencia e solicitacao LGPD  05 de outubro de 2026'); p.runs[0].font.size=Pt(11); p.runs[0].font.color.rgb=RGBColor.from_string('0C879E')

heading(doc,'Resultado')
text(doc,'O fluxo de formularios foi alterado para envio HTTPS ao endpoint /api/mail.php. O navegador nao monta mais uma URL mailto com dados pessoais. O endpoint aceita somente os dois tipos de formulario previstos, define o destinatario no servidor e devolve uma resposta controlada ao visitante.')

heading(doc,'Controles implementados')
table(doc,['Controle','Aplicacao'],[
 ('Destino fixo no servidor','Privacidade envia apenas para encarregado@life.eng.br e transparencia apenas para contato@life.eng.br.'),
 ('Validacao de origem','O endpoint aceita apenas POST JSON vindo da mesma origem hospedada.'),
 ('Validacao de dados','Campos permitidos, obrigatorios, e-mail e limites de tamanho sao conferidos no servidor.'),
 ('Antispam','Honeypot invisivel e limite de cinco envios por IP a cada hora.'),
 ('Prevencao de injecao','Quebras de linha sao removidas dos valores antes de montar assunto e cabecalhos.'),
 ('Privacidade','Nao ha armazenamento do conteudo no codigo; a resposta usa Cache Control no store.'),
])

heading(doc,'Mudancas no formulario')
text(doc,'Os formularios de Transparencia e Privacidade agora exibem uma mensagem de envio seguro e usam fetch com credenciais da mesma origem. Em caso de falha, o usuario recebe uma mensagem generica sem expor detalhes internos. Os enderecos de destino nao sao aceitos do navegador.')

heading(doc,'Requisito para a hospedagem UOL')
text(doc,'O arquivo deploy/uol/api/mail.php deve ser publicado como api/mail.php em um plano UOL com PHP e funcao mail habilitados. Ele nao deve ser colocado na pasta public de uma implantacao Next.js. O remetente no reply to vem apenas do campo de e-mail validado; o remetente tecnico e no-reply@life.eng.br. O dominio deve autorizar esse remetente no painel de e-mail ou no provedor responsavel por SPF e DKIM.')

heading(doc,'Validacao realizada')
table(doc,['Verificacao','Resultado'],[
 ('Build Next.js','Concluido com sucesso; 34 paginas e documentos de rota gerados.'),
 ('Fluxo do navegador','O codigo cliente envia somente JSON ao endpoint HTTPS da mesma origem.'),
 ('Validacao PHP','Pendente no ambiente UOL porque este ambiente local nao possui interpretador PHP.'),
 ('Entrega de e-mail','Pendente apos publicar; depende da funcao mail e da configuracao de remetente da UOL.'),
])

heading(doc,'Teste de aceite apos publicar')
text(doc,'Enviar uma solicitacao valida em cada formulario, confirmar recebimento no destinatario correto e resposta de sucesso na pagina. Em seguida, tentar campos obrigatorios ausentes, e-mail invalido, origem externa, honeypot preenchido e mais de cinco envios em uma hora. Os testes devem retornar bloqueio ou mensagem de erro sem revelar informacoes internas.')

f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=f.add_run('Life Engenharia  Relatorio tecnico do canal de email'); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string('667985')
doc.core_properties.title='Relatorio de seguranca do canal de email'; doc.core_properties.author='Life Engenharia'; doc.save(OUT)
print(OUT.resolve())
