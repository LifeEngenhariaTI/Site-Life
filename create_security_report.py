from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from pathlib import Path

OUT = Path('Relatorio_de_Vulnerabilidades_Life_Engenharia.docx')
NAVY = '07263D'
TEAL = '0C879E'
PALE = 'EEF5F7'
LINE = 'D9E3E8'

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def borders(cell, color=LINE):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in('w:tcBorders')
    if tc_borders is None:
        tc_borders = OxmlElement('w:tcBorders')
        tc_pr.append(tc_borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '6')
        element.set(qn('w:color'), color)

def set_cell(cell, text, bold=False, color=None, size=9):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(3)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = 'Aptos'
    r._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    borders(cell)

def para(doc, text='', bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead); r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    for r in p.runs:
        r.font.name = 'Aptos'; r._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos'); r.font.size = Pt(10.5)
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(15 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor.from_string('000000')
    return p

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]; shade(cell, NAVY); set_cell(cell, h, True, 'FFFFFF', 9)
        cell.width = Inches(widths[i])
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            if n % 2 == 1: shade(cells[i], PALE)
            set_cell(cells[i], value, size=8.7)
            cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return t

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.72); section.bottom_margin = Inches(.72)
section.left_margin = Inches(.78); section.right_margin = Inches(.78)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); styles['Normal'].font.size = Pt(10.5)
for name in ['Title', 'Heading 1', 'Heading 2']:
    styles[name].font.name = 'Aptos'; styles[name]._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); styles[name].font.color.rgb = RGBColor(0,0,0)
styles['Title'].font.size = Pt(24); styles['Heading 1'].font.size = Pt(15); styles['Heading 2'].font.size = Pt(12)

title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run('Relatorio de vulnerabilidades do site Life Engenharia')
subtitle = doc.add_paragraph('Revisao tecnica de codigo e configuracao local  05 de outubro de 2026')
subtitle.runs[0].font.size = Pt(11); subtitle.runs[0].font.color.rgb = RGBColor.from_string(TEAL)

heading(doc, 'Conclusao executiva')
para(doc, 'A revisao do projeto Next.js nao identificou vulnerabilidades conhecidas nas dependencias de producao. Foram identificados dois pontos de endurecimento de seguranca que devem ser resolvidos antes da publicacao: a ausencia de cabecalhos HTTP de seguranca na configuracao da aplicacao e o fluxo de formularios que transfere dados pessoais para o cliente de e-mail do visitante. Nenhuma exploracao ativa, varredura externa ou teste contra infraestrutura da UOL foi realizado.')

heading(doc, 'Escopo e evidencias')
table(doc, ['Item', 'Cobertura'], [
    ('Aplicacao', 'Codigo do projeto Next.js, paginas estaticas, assets e package lock.'),
    ('Dependencias', 'npm audit --omit=dev: 0 vulnerabilidades conhecidas em 16 dependencias de producao.'),
    ('Rotas', 'Build de producao concluido com 31 paginas pre renderizadas, sitemap.xml e robots.txt.'),
    ('Cabecalhos', 'Resposta local do Next.js revisada para a rota inicial; cabecalhos de endurecimento nao foram configurados no projeto.'),
    ('Fora do escopo', 'DNS, TLS, painel UOL, permissao de arquivos, WAF, e mail, banco de dados, logs, credenciais e testes autenticados.'),
], [1.35, 5.6])

heading(doc, 'Resumo dos achados')
table(doc, ['ID', 'Severidade', 'Achado', 'Prioridade'], [
    ('SEC 01', 'Media', 'Cabecalhos HTTP de seguranca ausentes na configuracao do Next.js.', 'Antes da publicacao'),
    ('SEC 02', 'Media', 'Solicitacoes de privacidade e contato usam mailto no navegador.', 'Antes da publicacao'),
    ('SEC 03', 'Baixa', 'Conteudo HTML local e injetado na renderizacao sem controle de integridade de conteudo.', 'Monitorar e controlar alteracoes'),
    ('SEC 04', 'Informativa', 'Nao ha endpoint de formulario no servidor, autenticacao ou banco de dados no codigo analisado.', 'Validar no ambiente hospedado'),
], [.65, .9, 3.7, 1.7])

heading(doc, 'Achados detalhados')
heading(doc, 'SEC 01 Cabecalhos HTTP de seguranca', 2)
para(doc, 'Severidade: Media. A configuracao next.config.mjs nao define Content Security Policy, X Content Type Options, Referrer Policy, Permissions Policy, X Frame Options ou Strict Transport Security. A ausencia desses controles reduz a protecao contra incorporacao indevida em frames, interpretacao de tipos de arquivo e carregamento de recursos nao previstos.')
para(doc, 'Recomendacao: configurar os cabecalhos no Next.js e confirmar que a hospedagem UOL preserva os cabecalhos na resposta publica. Iniciar com uma CSP em modo report only, ajustar origens necessarias para imagens e WhatsApp e depois aplicar a politica de bloqueio. HSTS deve ser ativado somente apos confirmar HTTPS funcional em todo o dominio e subdominios relevantes.')

heading(doc, 'SEC 02 Dados pessoais enviados por mailto', 2)
para(doc, 'Severidade: Media. Os formularios identificados por data mail form montam uma URL mailto no navegador. Os campos podem conter dados pessoais, e a entrega passa pelo cliente de e-mail e pela configuracao local do visitante. Isso impede controle consistente de entrega, limite de tamanho, rastreabilidade, retencao e tratamento de falhas.')
para(doc, 'Recomendacao: substituir mailto por um endpoint HTTPS no servidor, com validacao estrita por campo, limite de tamanho, protecao CSRF quando houver sessao, rate limiting, CAPTCHA ou desafio equivalente, registro minimo de auditoria e encaminhamento seguro ao encarregado. Definir tambem prazo de retencao e acesso aos dados conforme a politica de privacidade.')

heading(doc, 'SEC 03 Renderizacao de HTML local', 2)
para(doc, 'Severidade: Baixa. A rota coringa do App Router usa dangerouslySetInnerHTML para renderizar o HTML armazenado no repositorio. No estado atual o conteudo vem de arquivos locais, e nao de entrada do visitante. O risco aparece se pessoas sem revisao de codigo puderem alterar esses arquivos ou se o conteudo passar a vir de CMS, formulario ou fonte externa.')
para(doc, 'Recomendacao: manter revisao de alteracoes e permissao restrita de escrita nos arquivos de conteudo. Antes de adotar qualquer fonte externa, sanitizar HTML com lista de permissao e remover scripts, atributos de evento e URLs perigosas.')

heading(doc, 'SEC 04 Confirmacoes necessarias no ambiente hospedado', 2)
para(doc, 'Severidade: Informativa. O codigo revisado nao possui API de recebimento de dados, autenticacao, banco de dados ou segredos. A seguranca final depende tambem do ambiente UOL. Confirmar HTTPS com certificado valido, redirecionamento HTTP para HTTPS, permissoes de arquivos, atualizacao do runtime Node, backups, acesso ao painel e existencia de WAF ou firewall.')

heading(doc, 'Plano de correcao recomendado')
table(doc, ['Prazo', 'Acao', 'Criterio de aceite'], [
    ('Imediato', 'Adicionar cabecalhos de seguranca no Next.js e validar em HTTPS publico.', 'CSP, anti frame, nosniff, referrer e permissions policy retornam nas respostas.'),
    ('Imediato', 'Definir o destino do formulario como endpoint HTTPS ou manter aviso explicito de abertura do e-mail.', 'Nenhum dado pessoal e enviado a um destino nao documentado.'),
    ('Curto prazo', 'Adicionar limitacao de requisicoes e antispam ao endpoint de formulario.', 'Abuso automatizado e envios excessivos sao bloqueados e registrados.'),
    ('Curto prazo', 'Revisar configuracao da UOL e publicar checklist operacional.', 'HTTPS, backup, acesso administrativo e atualizacoes foram confirmados.'),
    ('Continuo', 'Executar npm audit e revisao de dependencias a cada atualizacao.', 'Nenhuma vulnerabilidade conhecida pendente em producao.'),
], [.8, 3.3, 2.5])

heading(doc, 'Limites da revisao')
para(doc, 'Este documento registra uma revisao defensiva do codigo local e de uma execucao local do Next.js. Ele nao certifica a infraestrutura de hospedagem nem substitui teste de penetracao autorizado. Achados ausentes neste relatorio nao devem ser interpretados como garantia de ausencia de vulnerabilidades.')

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('Life Engenharia  Relatorio tecnico de seguranca  05 de outubro de 2026')
run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string('667985')

doc.core_properties.title = 'Relatorio de vulnerabilidades do site Life Engenharia'
doc.core_properties.subject = 'Revisao tecnica de seguranca'
doc.core_properties.author = 'Life Engenharia'
doc.save(OUT)
print(OUT.resolve())
